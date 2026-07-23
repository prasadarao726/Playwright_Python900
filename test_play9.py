from pathlib import Path

import pytest
from openpyxl import load_workbook


@pytest.mark.excel999
def test_excelhandling():
    excel_path = Path(__file__).resolve().parents[1] / "testdata" / "123excel.xlsx"
    exceldata = load_workbook(excel_path)
    sheet = exceldata["Sheet1"]

    values = [row for row in sheet.iter_rows(min_row=2, values_only=True)]

    print(values)
    assert values, "No data found in the Excel sheet"
    print(values[0])
    print(values[1])